"""Read sales and stock files exported from selling platforms.

Sellers upload whatever their platform gives them (eBay, Amazon, Shopify, Etsy,
TikTok Shop, or their own spreadsheet). This module:

  1. reads the file (CSV, tab-separated .txt, or Excel .xlsx), in any common
     encoding and delimiter, skipping title rows above the real header;
  2. recognises the platform from its column names;
  3. works out which column holds what ("Item Title" -> product name, ...);
  4. turns rows into clean order lines or product/stock records, leaving out
     cancelled and refunded orders.

Everything here is pure (no database); `services_import` applies the result.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from typing import Iterable

MAX_ROWS = 100_000

# ---------------------------------------------------------------------------------------------
# Column roles and the header names that mean them (compared after normalising: lower case,
# letters and digits only). Order matters: earlier names win when a file has several.
# ---------------------------------------------------------------------------------------------

ROLE_SYNONYMS: dict[str, list[str]] = {
    "date": [
        "saledate", "purchasedate", "createdat", "createdtime", "orderdate", "datepaid", "paidondate", "paidat",
        "paidtime", "orderedon", "date", "solddate", "orderplaceddate", "ordercreatedtime", "transactiondate", "dateordered",
    ],
    "order_id": [
        "ordernumber", "amazonorderid", "orderid", "salesrecordnumber", "orderno", "order",
        "orderref", "orderreference", "transactionid",
    ],
    "sku": [
        "customlabel", "sellersku", "lineitemsku", "variantsku", "sku", "skuid", "itemsku", "productsku", "merchantsku",
        "partnumber", "productcode", "itemcode", "code",
    ],
    "item_id": ["itemnumber", "asin", "listingid", "productid", "itemid", "ebayitemid"],
    "name": [
        "itemtitle", "lineitemname", "productname", "itemname", "title", "product", "productnames", "products",
        "item", "description", "listingtitle", "name",
    ],
    "quantity": [
        "quantity", "lineitemquantity", "qty", "quantitypurchased", "quantityordered", "itemssold", "unitssold",
        "units", "quantitysold", "soldquantity", "qtysold",
    ],
    "unit_price": [
        "soldfor", "lineitemprice", "price", "skuunitoriginalprice", "unitprice", "itemprice", "saleprice",
        "priceeach", "sellingprice", "pricepaid", "currentprice", "startprice", "buyitnowprice", "regularprice",
        "yourprice", "retailprice",
    ],
    "line_total": [
        "itemtotal", "skusubtotalafterdiscount", "skusubtotalbeforediscount", "linetotal", "netsales",
        "total", "totalprice", "amount", "revenue", "sales",
    ],
    "status": ["orderstatus", "status", "itemstatus", "financialstatus", "orderstate"],
    "channel": ["saleschannel", "source", "channel", "marketplace", "platform"],
    "stock": [
        "variantinventoryqty", "availablequantity", "quantityavailable", "stock", "stockqty", "instock", "onhand",
        "inventory", "inventoryqty", "afnfulfillablequantity", "fulfillablequantity", "available", "stocklevel",
        "stockquantity", "qtyonhand",
    ],
    "cost": ["costperitem", "cost", "unitcost", "costprice", "buyprice", "purchaseprice", "suppliercost", "landedcost"],
    "category": ["type", "category", "productcategory", "producttype", "categoryname"],
    "supplier": ["vendor", "supplier", "brand", "manufacturer"],
    "lead_time_days": ["leadtimedays", "leadtime", "deliverydays", "shippingdays", "supplierdays"],
}

ROLE_LABELS = {
    "date": "Sale date",
    "order_id": "Order number",
    "sku": "SKU / custom label",
    "item_id": "Listing / item ID",
    "name": "Product name",
    "quantity": "Quantity sold",
    "unit_price": "Price per item",
    "line_total": "Line total",
    "status": "Order status",
    "channel": "Sales channel",
    "stock": "Stock on hand",
    "cost": "Cost per item",
    "category": "Category",
    "supplier": "Supplier / brand",
    "lead_time_days": "Supplier lead time (days)",
}

# Header fingerprints. Each entry: platform -> (kind, normalised headers that must all be present).
PLATFORM_SIGNATURES: list[tuple[str, str, set[str]]] = [
    ("amazon", "orders", {"amazonorderid", "purchasedate"}),
    ("amazon", "stock", {"sellersku", "afnfulfillablequantity"}),
    ("ebay", "orders", {"salesrecordnumber", "itemtitle"}),
    ("ebay", "orders", {"salesrecordnumber", "soldfor"}),
    ("ebay", "stock", {"itemnumber", "availablequantity"}),
    ("shopify", "orders", {"lineitemquantity", "lineitemname"}),
    ("shopify", "stock", {"handle", "variantsku"}),
    ("shopify", "stock", {"handle", "variantinventoryqty"}),
    ("etsy", "orders", {"saledate", "itemname", "listingid"}),
    ("etsy", "orders", {"saledate", "itemname", "transactionid"}),
    ("tiktok", "orders", {"sellersku", "createdtime"}),
    ("tiktok", "orders", {"skuunitoriginalprice"}),
    ("woocommerce", "orders", {"itemssold", "netsales"}),
]

PLATFORM_NAMES = {
    "ebay": "eBay",
    "amazon": "Amazon",
    "shopify": "Shopify",
    "etsy": "Etsy",
    "tiktok": "TikTok Shop",
    "woocommerce": "WooCommerce",
    "spreadsheet": "your spreadsheet",
}

# Amazon's "item-price" is the price for the whole line (all units), not per unit.
LINE_TOTAL_PRICE_PLATFORMS = {"amazon"}

SKIP_STATUS = re.compile(r"cancel|refund|return|void|fail|declin|unpaid|pending payment", re.I)


def norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (header or "").lower())


# ---------------------------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------------------------


class ImportError_(Exception):
    """A problem the seller can fix; the message is shown to them as is."""


def read_tables(filename: str, data: bytes) -> list[tuple[str, list[list[str]]]]:
    """Return [(sheet name, rows of strings)]: one entry for CSV/TXT, one per useful Excel sheet."""
    name = (filename or "").lower()
    if not data:
        raise ImportError_("The file is empty.")
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return _read_xlsx(data)
    if name.endswith(".xls") or data[:4] == b"\xd0\xcf\x11\xe0":
        raise ImportError_(
            "This is an old Excel file (.xls). Open it in Excel or Google Sheets and choose "
            "File > Save as > .xlsx or CSV, then upload that."
        )
    if name.endswith((".numbers", ".pdf", ".zip")):
        raise ImportError_("Please upload a CSV, TXT or Excel (.xlsx) file.")
    text = _decode(data)
    sample = text[:20000]
    delim = _sniff_delimiter(sample)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return [("", [[c.strip() for c in r] for r in rows if any(c.strip() for c in r)][:MAX_ROWS])]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _sniff_delimiter(sample: str) -> str:
    lines = [ln for ln in sample.splitlines()[:30] if ln.strip()]
    best, best_score = ",", -1
    for d in (",", "\t", ";", "|"):
        counts = [ln.count(d) for ln in lines]
        if not counts:
            continue
        # Prefer the delimiter that appears often and consistently on many lines.
        common = max(set(counts), key=counts.count)
        score = common * counts.count(common) if common else 0
        if score > best_score:
            best, best_score = d, score
    return best


def _read_xlsx(data: bytes) -> list[tuple[str, list[list[str]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise ImportError_("Excel files aren't supported on this server yet. Please upload a CSV.") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any corrupt workbook
        raise ImportError_("We couldn't open this Excel file. Try saving it again as .xlsx or CSV.") from exc
    sheets: list[tuple[str, list[list[str]]]] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            if any(cells):
                rows.append(cells)
            if len(rows) >= MAX_ROWS:
                break
        # Keep sheets whose header names at least two columns we understand.
        if _header_score(rows) >= 2:
            sheets.append((ws.title, rows))
    if not sheets:
        raise ImportError_(
            "We couldn't find sales or stock columns in this Excel file. Make sure the first rows of a sheet "
            "have headings like Date, Product, Quantity, Price or Stock."
        )
    return sheets


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _header_score(rows: list[list[str]]) -> int:
    if not rows:
        return -1
    return max(_row_role_hits(r) for r in rows[:15])


def _row_role_hits(row: list[str]) -> int:
    known = {s for syns in ROLE_SYNONYMS.values() for s in syns}
    return sum(1 for c in row if norm(c) in known)


def find_header(rows: list[list[str]]) -> int:
    """Index of the header row: the row near the top naming the most known columns.

    eBay and some Excel sheets put a title or blank lines above the real header.
    """
    best_i, best_hits = 0, -1
    for i, r in enumerate(rows[:15]):
        hits = _row_role_hits(r)
        if hits > best_hits:
            best_i, best_hits = i, hits
    return best_i


# ---------------------------------------------------------------------------------------------
# Understanding the columns
# ---------------------------------------------------------------------------------------------


@dataclass
class Analysis:
    platform: str  # key of PLATFORM_NAMES
    kind: str  # "orders" | "stock"
    header_row: int
    columns: list[str]
    mapping: dict[str, int]  # role -> column index
    missing: list[str]  # required roles we could not find
    confident: bool
    sample_rows: list[list[str]] = field(default_factory=list)

    def describe(self) -> list[dict]:
        """Plain-language list of what we read from which column."""
        return [
            {"role": role, "label": ROLE_LABELS[role], "column": self.columns[idx]}
            for role, idx in sorted(self.mapping.items(), key=lambda kv: kv[1])
        ]


def detect_platform(headers: list[str]) -> tuple[str, str | None]:
    have = {norm(h) for h in headers}
    for platform, kind, needed in PLATFORM_SIGNATURES:
        if needed <= have:
            return platform, kind
    return "spreadsheet", None


# Platform-specific meanings that would be wrong in a general spreadsheet.
PLATFORM_OVERRIDES: dict[str, dict[str, list[str]]] = {
    "shopify": {"order_id": ["name"]},  # Shopify's "Name" column is the order number (#1001)
}

# Month-first dates when a file's dates are ambiguous (every day number is 12 or lower).
MONTH_FIRST_PLATFORMS = {"etsy", "amazon", "tiktok", "shopify"}


def map_columns(headers: list[str], platform: str = "spreadsheet") -> dict[str, int]:
    normed = [norm(h) for h in headers]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for role, synonyms in PLATFORM_OVERRIDES.get(platform, {}).items():
        idx = next((i for i, h in enumerate(normed) if h in synonyms), None)
        if idx is not None:
            mapping[role] = idx
            used.add(idx)
    for role, synonyms in ROLE_SYNONYMS.items():
        if role in mapping:
            continue
        for syn in synonyms:
            idx = next((i for i, h in enumerate(normed) if h == syn and i not in used), None)
            if idx is not None:
                mapping[role] = idx
                used.add(idx)
                break
    # Loose second pass for headers like "Quantity (units)" or "Sale date (UTC)".
    for role, synonyms in ROLE_SYNONYMS.items():
        if role in mapping:
            continue
        for syn in synonyms:
            if len(syn) < 4:
                continue
            idx = next((i for i, h in enumerate(normed) if i not in used and h.startswith(syn)), None)
            if idx is not None:
                mapping[role] = idx
                used.add(idx)
                break
    return mapping


def required_roles(kind: str) -> list[str]:
    if kind == "orders":
        return ["date", "quantity"]
    return ["stock"]


def analyse(rows: list[list[str]]) -> Analysis:
    if not rows:
        raise ImportError_("The file is empty.")
    h = find_header(rows)
    headers = rows[h]
    platform, kind = detect_platform(headers)
    mapping = map_columns(headers, platform)
    if platform == "shopify":
        mapping.pop("line_total", None)  # Shopify's "Total" is the whole order, not one line
    if kind is None:
        kind = "orders" if "date" in mapping and "quantity" in mapping else "stock"
    # A product must be identifiable by SKU, item ID or name.
    missing = [r for r in required_roles(kind) if r not in mapping]
    if not any(r in mapping for r in ("sku", "item_id", "name")):
        missing.append("name")
    confident = not missing and (platform != "spreadsheet" or len(mapping) >= 3)
    return Analysis(platform, kind, h, headers, mapping, missing, confident, rows[h + 1 : h + 6])


# ---------------------------------------------------------------------------------------------
# Turning rows into records
# ---------------------------------------------------------------------------------------------

_DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
    "%m/%d/%y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y", "%b %d, %Y", "%d %b %Y", "%b-%d-%y", "%d-%b-%y",
    "%B %d, %Y", "%d %B %Y", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p", "%d/%m/%Y %I:%M:%S %p",
]

_TZ_SUFFIX = re.compile(r"\s*(PST|PDT|EST|EDT|CST|CDT|MST|MDT|GMT|UTC|BST|CET|CEST|AEST|AEDT)$", re.I)


def parse_date(value: str, day_first: bool | None = None) -> datetime | None:
    v = (value or "").strip()
    if not v:
        return None
    v = _TZ_SUFFIX.sub("", v)
    try:
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
        return dt.astimezone(timezone.utc) if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    except ValueError:
        pass
    v = re.sub(r"\s+[+-]\d{2}:?\d{2}$", "", v)  # strip "+0100" style offsets
    formats = _DATE_FORMATS
    if day_first is False:
        formats = [f for f in formats if not f.startswith("%d/")]
    elif day_first is True:
        formats = [f for f in formats if not f.startswith("%m/")]
    for fmt in formats:
        try:
            return datetime.strptime(v, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def guess_day_first(values: Iterable[str]) -> bool | None:
    """Decide between 03/04/2026 = 3 April (UK) or 4 March (US) from the whole column."""
    first_over_12 = second_over_12 = False
    for v in values:
        m = re.match(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-]\d{2,4}", v or "")
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        first_over_12 |= a > 12
        second_over_12 |= b > 12
    if first_over_12 and not second_over_12:
        return True
    if second_over_12 and not first_over_12:
        return False
    return None


_NUM = re.compile(r"-?\d+(?:[.,]\d+)*")


def parse_number(value: str) -> float | None:
    v = (value or "").strip()
    if not v:
        return None
    neg = v.startswith("(") and v.endswith(")")
    m = _NUM.search(v.replace(" ", "").replace(" ", ""))
    if not m:
        return None
    s = m.group(0)
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        # "1,234" (thousands) vs "12,50" (decimal comma)
        s = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def _slug(text: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "-", text.upper()).strip("-")[:40] or "ITEM"


@dataclass
class OrderLine:
    key: str  # product identity: SKU, else item id, else name
    sku: str
    name: str
    quantity: int
    unit_price: float | None
    sold_at: datetime
    order_ref: str
    channel: str
    import_key: str


@dataclass
class StockRecord:
    key: str
    sku: str
    name: str
    stock: int | None
    price: float | None
    cost: float | None
    category: str
    supplier: str
    lead_time_days: int | None


@dataclass
class Extracted:
    orders: list[OrderLine] = field(default_factory=list)
    stock: list[StockRecord] = field(default_factory=list)
    skipped_cancelled: int = 0
    skipped_invalid: int = 0
    problems: list[str] = field(default_factory=list)
    first_date: datetime | None = None
    last_date: datetime | None = None


def _get(row: list[str], mapping: dict[str, int], role: str) -> str:
    idx = mapping.get(role)
    return row[idx].strip() if idx is not None and idx < len(row) else ""


def product_key(sku: str, item_id: str, name: str, platform: str) -> tuple[str, str]:
    """Return (key, sku). Rows without a SKU get one made from the listing ID or the name."""
    if sku:
        return sku.upper(), sku
    if item_id:
        made = f"{platform.upper()[:6]}-{item_id}"
        return made.upper(), made
    if name:
        made = _slug(name)
        return made, made
    return "", ""


def extract(rows: list[list[str]], a: Analysis) -> Extracted:
    out = Extracted()
    body = rows[a.header_row + 1 :]
    m = a.mapping
    if a.kind == "orders":
        day_first = guess_day_first(_get(r, m, "date") for r in body[:2000])
        if day_first is None and a.platform in MONTH_FIRST_PLATFORMS:
            day_first = False
        cancelled_col = next((i for i, h in enumerate(a.columns) if norm(h) in ("cancelledat", "cancelleddate")), None)
        last_order = ""
        last_date_raw = ""
        cancelled_orders: set[str] = set()
        seen: dict[str, int] = {}
        for i, r in enumerate(body, start=a.header_row + 2):
            if _row_is_footer(r):
                continue
            # Shopify fills order-level columns (date, status) only on an order's first row.
            order_ref = _get(r, m, "order_id") or last_order
            raw_date = _get(r, m, "date") or (last_date_raw if order_ref == last_order else "")
            last_order, last_date_raw = order_ref, raw_date
            status = _get(r, m, "status")
            cancelled = cancelled_col is not None and cancelled_col < len(r) and r[cancelled_col].strip()
            if cancelled or (status and SKIP_STATUS.search(status)) or (order_ref and order_ref in cancelled_orders):
                if order_ref:
                    cancelled_orders.add(order_ref)
                out.skipped_cancelled += 1
                continue
            name = _get(r, m, "name")
            key, sku = product_key(_get(r, m, "sku"), _get(r, m, "item_id"), name, a.platform)
            qty_raw = _get(r, m, "quantity")
            qty = parse_number(qty_raw) if qty_raw else 1.0
            when = parse_date(raw_date, day_first)
            if not key or when is None or qty is None or qty <= 0:
                out.skipped_invalid += 1
                if len(out.problems) < 20:
                    reason = "no product name or SKU" if not key else "no readable date" if when is None else "no quantity"
                    out.problems.append(f"Row {i}: {reason}")
                continue
            quantity = int(round(qty))
            price = parse_number(_get(r, m, "unit_price"))
            total = parse_number(_get(r, m, "line_total"))
            if a.platform in LINE_TOTAL_PRICE_PLATFORMS and price is not None:
                price = price / quantity
            if price is None and total is not None:
                price = total / quantity
            channel = _get(r, m, "channel") or PLATFORM_NAMES.get(a.platform, "")
            if a.platform == "spreadsheet" and channel == PLATFORM_NAMES["spreadsheet"]:
                channel = ""
            ident = "|".join([a.platform, order_ref, key, str(quantity), when.isoformat(), str(price)])
            # Identical lines can legitimately repeat inside one file; number them so each is kept,
            # while uploading the same file again still produces the same keys.
            seen[ident] = seen.get(ident, 0) + 1
            ident += f"#{seen[ident]}"
            out.orders.append(
                OrderLine(
                    key=key, sku=sku, name=name or sku, quantity=quantity, unit_price=price, sold_at=when,
                    order_ref=order_ref, channel=channel.lower(), import_key=hashlib.sha1(ident.encode()).hexdigest()[:24],
                )
            )
            out.first_date = when if out.first_date is None or when < out.first_date else out.first_date
            out.last_date = when if out.last_date is None or when > out.last_date else out.last_date
    else:
        for i, r in enumerate(body, start=a.header_row + 2):
            if _row_is_footer(r):
                continue
            name = _get(r, m, "name")
            key, sku = product_key(_get(r, m, "sku"), _get(r, m, "item_id"), name, a.platform)
            if not key:
                out.skipped_invalid += 1
                continue
            stock = parse_number(_get(r, m, "stock"))
            lead = parse_number(_get(r, m, "lead_time_days"))
            out.stock.append(
                StockRecord(
                    key=key, sku=sku, name=name or sku,
                    stock=int(round(stock)) if stock is not None else None,
                    price=parse_number(_get(r, m, "unit_price")),
                    cost=parse_number(_get(r, m, "cost")),
                    category=_get(r, m, "category"), supplier=_get(r, m, "supplier"),
                    lead_time_days=int(lead) if lead is not None and 0 <= lead <= 180 else None,
                )
            )
        # Shopify lists one row per variant/image; keep the last row that carries data per key.
        merged: dict[str, StockRecord] = {}
        for rec in out.stock:
            prev = merged.get(rec.key)
            if prev is None:
                merged[rec.key] = rec
            else:
                for f in ("stock", "price", "cost", "lead_time_days"):
                    if getattr(rec, f) is not None:
                        setattr(prev, f, getattr(rec, f))
                for f in ("category", "supplier"):
                    if getattr(rec, f):
                        setattr(prev, f, getattr(rec, f))
                if rec.name and (not prev.name or prev.name == prev.sku):
                    prev.name = rec.name
        out.stock = list(merged.values())
    return out


def _row_is_footer(r: list[str]) -> bool:
    joined = " ".join(r).lower()
    return bool(re.search(r"record\(s\) downloaded|seller id :|^total\b|^totals\b", joined)) and sum(1 for c in r if c) <= 3


def summary(a: Analysis, x: Extracted) -> dict:
    """Plain-language description of what an import will do, for the confirm screen."""
    platform = PLATFORM_NAMES.get(a.platform, a.platform)
    if a.kind == "orders":
        products = {o.key for o in x.orders}
        units = sum(o.quantity for o in x.orders)
        span = ""
        if x.first_date and x.last_date:
            span = f" from {x.first_date:%d %b %Y} to {x.last_date:%d %b %Y}"
        headline = f"We found {units:,} items sold across {len(products):,} products{span} in your {platform} file."
        if a.platform == "spreadsheet":
            headline = f"We found {units:,} items sold across {len(products):,} products{span}."
    else:
        headline = f"We found {len(x.stock):,} products with stock levels in your {platform} file."
        if a.platform == "spreadsheet":
            headline = f"We found {len(x.stock):,} products with stock levels."
    notes = []
    if x.skipped_cancelled:
        notes.append(f"{x.skipped_cancelled:,} cancelled or refunded orders will be left out.")
    if x.skipped_invalid:
        notes.append(f"{x.skipped_invalid:,} rows will be skipped because they were incomplete.")
    if a.kind == "orders" and x.last_date and x.last_date < datetime.now(timezone.utc) - timedelta(days=21):
        notes.append("The newest sale in this file is over 3 weeks old. Download a file that includes recent sales for better forecasts.")
    if a.kind == "orders" and x.first_date and x.last_date and (x.last_date - x.first_date).days < 42:
        notes.append("This file covers less than 6 weeks. Forecasts get much better with 3 months or more of sales.")
    return {"headline": headline, "notes": notes}
